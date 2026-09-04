import os
import matplotlib.pyplot as plt
import matplotlib.patches as patches
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch, Circle
from pptx import Presentation
from pptx.util import Inches, Pt
import glob

# Set Chinese font
plt.rcParams['font.sans-serif'] = ['Microsoft YaHei', 'SimHei', 'Arial']
plt.rcParams['axes.unicode_minus'] = False

out_dir = r"d:\SmartLab2.0\doc\generated_images"
os.makedirs(out_dir, exist_ok=True)

# =============================================================
# 1. 纵向概念结构图 (单步闭环)
# =============================================================
def create_vertical_diagram():
    fig, ax = plt.subplots(figsize=(8.2, 7.2), dpi=300)
    fig.patch.set_facecolor('#F8FAFC')
    ax.set_facecolor('#F8FAFC')
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis('off')

    # Outer Card Border
    outer_card = FancyBboxPatch((1.5, 1.5), 97, 97, boxstyle="round,pad=0.5,rounding_size=2.0",
                                facecolor='#FFFFFF', edgecolor='#CBD5E1', linewidth=1.5, zorder=1)
    ax.add_patch(outer_card)

    # Header Title
    ax.text(50, 94.2, "纵向概念结构图 (单步执行闭环)", ha='center', va='center',
            fontsize=15.5, fontweight='bold', color='#0F172A', zorder=4)
    ax.text(50, 90.2, "单步控制必穿透状态机 · 严禁旁路直连 · 反馈原路严格收敛", ha='center', va='center',
            fontsize=9.2, color='#64748B', zorder=4)

    # 1. Workflow Node
    node1 = FancyBboxPatch((18, 74.0), 64, 11.5, boxstyle="round,pad=0.5,rounding_size=1.8",
                           facecolor='#1E3A8A', edgecolor='#3B82F6', linewidth=2.2, zorder=3)
    ax.add_patch(node1)
    ax.text(50, 81.2, "工作流执行节点 (Workflow Node)", ha='center', va='center',
            fontsize=12, fontweight='bold', color='#FFFFFF', zorder=4)
    ax.text(50, 76.8, "DEV_NODE · 声明能力映射 · 维护步骤生命周期", ha='center', va='center',
            fontsize=8.5, color='#93C5FD', zorder=4)

    # Down Arrow 1
    a_down1 = FancyArrowPatch((35, 74.0), (35, 60.5),
                              arrowstyle='-|>', mutation_scale=14, linewidth=2.2,
                              color='#2563EB', zorder=5)
    ax.add_patch(a_down1)
    ax.text(33, 67.2, "下行控制\nWF_START", ha='right', va='center',
            fontsize=8.5, fontweight='bold', color='#1D4ED8', zorder=6)

    # Up Arrow 1
    a_up1 = FancyArrowPatch((65, 60.5), (65, 74.0),
                            arrowstyle='-|>', mutation_scale=14, linewidth=2.2,
                            linestyle='--', color='#059669', zorder=5)
    ax.add_patch(a_up1)
    ax.text(67, 67.2, "状态反馈\nCMD_STATE", ha='left', va='center',
            fontsize=8.5, fontweight='bold', color='#047857', zorder=6)

    # 2. State Machine (Core Highlight)
    sm_box = FancyBboxPatch((13, 42.0), 74, 17.5, boxstyle="round,pad=0.6,rounding_size=2.2",
                            facecolor='#F5F3FF', edgecolor='#7C3AED', linewidth=2.4, zorder=3)
    ax.add_patch(sm_box)
    
    # State machine badge
    badge = FancyBboxPatch((31, 57.2), 38, 4.0, boxstyle="round,pad=0.2,rounding_size=1.0",
                           facecolor='#7C3AED', edgecolor='#6D28D9', linewidth=1, zorder=4)
    ax.add_patch(badge)
    ax.text(50, 59.2, "★ 发令必经之路 · 独占硬件控制权", ha='center', va='center',
            fontsize=8.5, fontweight='bold', color='#FFFFFF', zorder=5)

    ax.text(50, 52.8, "设备状态机 (Device State Machine)", ha='center', va='center',
            fontsize=12.5, fontweight='bold', color='#5B21B6', zorder=4)

    # Sub compartments inside State Machine
    sub1 = FancyBboxPatch((16.5, 43.8), 31.5, 5.8, boxstyle="round,pad=0.2,rounding_size=0.8",
                          facecolor='#EDE9FE', edgecolor='#8B5CF6', linewidth=1, zorder=4)
    ax.add_patch(sub1)
    ax.text(32.25, 46.7, "CMD 空间: IDLE→SENT→RUNNING", ha='center', va='center',
            fontsize=7.8, fontweight='bold', color='#4C1D95', zorder=5)

    sub2 = FancyBboxPatch((52, 43.8), 31.5, 5.8, boxstyle="round,pad=0.2,rounding_size=0.8",
                          facecolor='#EDE9FE', edgecolor='#8B5CF6', linewidth=1, zorder=4)
    ax.add_patch(sub2)
    ax.text(67.75, 46.7, "OP 空间: OPERATIONAL / 异常锁存", ha='center', va='center',
            fontsize=7.8, fontweight='bold', color='#4C1D95', zorder=5)

    # Down Arrow 2
    a_down2 = FancyArrowPatch((35, 42.0), (35, 30.0),
                              arrowstyle='-|>', mutation_scale=14, linewidth=2.2,
                              color='#7C3AED', zorder=5)
    ax.add_patch(a_down2)
    ax.text(33, 36.0, "指令下发\nCMD_START", ha='right', va='center',
            fontsize=8.5, fontweight='bold', color='#6D28D9', zorder=6)

    # Up Arrow 2
    a_up2 = FancyArrowPatch((65, 30.0), (65, 42.0),
                            arrowstyle='-|>', mutation_scale=14, linewidth=2.2,
                            linestyle='--', color='#059669', zorder=5)
    ax.add_patch(a_up2)
    ax.text(67, 36.0, "事件回传\nevent / telemetry", ha='left', va='center',
            fontsize=8.5, fontweight='bold', color='#047857', zorder=6)

    # 3. Edge Adapter
    node3 = FancyBboxPatch((18, 18.5), 64, 11.0, boxstyle="round,pad=0.5,rounding_size=1.8",
                           facecolor='#0F766E', edgecolor='#14B8A6', linewidth=2.2, zorder=3)
    ax.add_patch(node3)
    ax.text(50, 25.2, "边缘适配器 (Edge Adapter)", ha='center', va='center',
            fontsize=12, fontweight='bold', color='#FFFFFF', zorder=4)
    ax.text(50, 21.0, "四层架构(runtime/NB/core/SB) · internal参数注入", ha='center', va='center',
            fontsize=8.5, color='#99F6E4', zorder=4)

    # Down Arrow 3
    a_down3 = FancyArrowPatch((35, 18.5), (35, 12.0),
                              arrowstyle='-|>', mutation_scale=13, linewidth=1.8,
                              color='#0D9488', zorder=5)
    ax.add_patch(a_down3)
    ax.text(33, 15.2, "物理控制帧", ha='right', va='center',
            fontsize=7.8, color='#0F766E', zorder=6)

    # Up Arrow 3
    a_up3 = FancyArrowPatch((65, 12.0), (65, 18.5),
                            arrowstyle='-|>', mutation_scale=13, linewidth=1.8,
                            linestyle='--', color='#059669', zorder=5)
    ax.add_patch(a_up3)
    ax.text(67, 15.2, "物理采样回传", ha='left', va='center',
            fontsize=7.8, color='#047857', zorder=6)

    # 4. Physical Equipment
    node4 = FancyBboxPatch((22, 4.5), 56, 7.5, boxstyle="round,pad=0.4,rounding_size=1.5",
                           facecolor='#1E293B', edgecolor='#475569', linewidth=1.8, zorder=3)
    ax.add_patch(node4)
    ax.text(50, 9.3, "物理仪器装备 / PLC", ha='center', va='center',
            fontsize=10.5, fontweight='bold', color='#F8FAFC', zorder=4)
    ax.text(50, 6.2, "反应釜 · 加热搅拌单元 · 传感器阵列", ha='center', va='center',
            fontsize=7.5, color='#94A3B8', zorder=4)

    out_path = os.path.join(out_dir, "01_纵向概念结构图.png")
    plt.tight_layout()
    plt.savefig(out_path, dpi=300, facecolor=fig.get_facecolor(), edgecolor='none')
    plt.close()
    print("Saved:", out_path)


# =============================================================
# 2. 横向概念结构图 (多步时序编排)
# =============================================================
def create_horizontal_diagram():
    fig, ax = plt.subplots(figsize=(8.2, 7.2), dpi=300)
    fig.patch.set_facecolor('#F8FAFC')
    ax.set_facecolor('#F8FAFC')
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis('off')

    # Outer Card Border
    outer_card = FancyBboxPatch((1.5, 1.5), 97, 97, boxstyle="round,pad=0.5,rounding_size=2.0",
                                facecolor='#FFFFFF', edgecolor='#CBD5E1', linewidth=1.5, zorder=1)
    ax.add_patch(outer_card)

    # Header Title
    ax.text(50, 94.2, "横向概念结构图 (多步时序编排)", ha='center', va='center',
            fontsize=15.5, fontweight='bold', color='#0F172A', zorder=4)
    ax.text(50, 90.2, "多步时序显式写出 · 控制流与数据流解耦 · 替代人工跨机流转", ha='center', va='center',
            fontsize=9.2, color='#64748B', zorder=4)

    # Time arrow banner
    t_arrow = FancyArrowPatch((12, 83.5), (88, 83.5), arrowstyle='-|>', mutation_scale=12,
                              linewidth=1.5, color='#94A3B8', zorder=2)
    ax.add_patch(t_arrow)
    ax.text(50, 85.5, "任务执行时序推移方向 (Time Axis)  t →", ha='center', va='center',
            fontsize=8.2, fontweight='bold', color='#64748B', zorder=3)

    # Workflow Container Frame
    wf_frame = FancyBboxPatch((4, 18), 92, 62, boxstyle="round,pad=0.6,rounding_size=2.0",
                              facecolor='#F8FAFC', edgecolor='#94A3B8', linewidth=1.5, linestyle='--', zorder=2)
    ax.add_patch(wf_frame)
    ax.text(7, 76.5, "任务工作流定义 (Workflow Definition)", fontsize=9.2, fontweight='bold', color='#334155', zorder=3)

    # 1. START Terminal
    start_c = patches.Circle((10, 52), 4.2, facecolor='#10B981', edgecolor='#059669', linewidth=2, zorder=4)
    ax.add_patch(start_c)
    ax.text(10, 52, "START", ha='center', va='center', fontsize=7.2, fontweight='bold', color='#FFFFFF', zorder=5)

    # Step 1: 加样
    s1_box = FancyBboxPatch((20, 36), 19, 32, boxstyle="round,pad=0.4,rounding_size=1.5",
                            facecolor='#FFFFFF', edgecolor='#2563EB', linewidth=2, zorder=4)
    ax.add_patch(s1_box)
    badge1 = FancyBboxPatch((22.5, 62.5), 14, 3.2, boxstyle="round,pad=0.1,rounding_size=0.6",
                            facecolor='#DBEAFE', edgecolor='#3B82F6', linewidth=0.8, zorder=5)
    ax.add_patch(badge1)
    ax.text(29.5, 64.1, "DEV_NODE", ha='center', va='center', fontsize=7.0, fontweight='bold', color='#1D4ED8', zorder=6)
    ax.text(29.5, 56.5, "步骤 1\n固体加样", ha='center', va='center', fontsize=11, fontweight='bold', color='#0F172A', zorder=5)
    ax.text(29.5, 48.0, "执行能力: dispense\n目标设定: 50.0g\n设备: 自动分配仪", ha='center', va='center', fontsize=7.5, color='#475569', zorder=5)
    
    # Port out Step 1
    p_out1 = FancyBboxPatch((22, 38), 15, 3.2, boxstyle="round,pad=0.1,rounding_size=0.4",
                            facecolor='#FFEDD5', edgecolor='#F97316', linewidth=0.8, zorder=5)
    ax.add_patch(p_out1)
    ax.text(29.5, 39.6, "PORT_OUT: 样品批次", ha='center', va='center', fontsize=6.5, fontweight='bold', color='#C2410C', zorder=6)

    # Step 2: 加热
    s2_box = FancyBboxPatch((46, 36), 19, 32, boxstyle="round,pad=0.4,rounding_size=1.5",
                            facecolor='#FFFFFF', edgecolor='#2563EB', linewidth=2, zorder=4)
    ax.add_patch(s2_box)
    badge2 = FancyBboxPatch((48.5, 62.5), 14, 3.2, boxstyle="round,pad=0.1,rounding_size=0.6",
                            facecolor='#DBEAFE', edgecolor='#3B82F6', linewidth=0.8, zorder=5)
    ax.add_patch(badge2)
    ax.text(55.5, 64.1, "DEV_NODE", ha='center', va='center', fontsize=7.0, fontweight='bold', color='#1D4ED8', zorder=6)
    ax.text(55.5, 56.5, "步骤 2\n恒温反应", ha='center', va='center', fontsize=11, fontweight='bold', color='#0F172A', zorder=5)
    ax.text(55.5, 48.0, "执行能力: heat\n目标设定: 80℃\n设备: 夹套反应釜", ha='center', va='center', fontsize=7.5, color='#475569', zorder=5)

    # Port out Step 2
    p_out2 = FancyBboxPatch((48, 38), 15, 3.2, boxstyle="round,pad=0.1,rounding_size=0.4",
                            facecolor='#FFEDD5', edgecolor='#F97316', linewidth=0.8, zorder=5)
    ax.add_patch(p_out2)
    ax.text(55.5, 39.6, "PORT_OUT: 反应状态", ha='center', va='center', fontsize=6.5, fontweight='bold', color='#C2410C', zorder=6)

    # Step 3: 搅拌
    s3_box = FancyBboxPatch((72, 36), 19, 32, boxstyle="round,pad=0.4,rounding_size=1.5",
                            facecolor='#FFFFFF', edgecolor='#2563EB', linewidth=2, zorder=4)
    ax.add_patch(s3_box)
    badge3 = FancyBboxPatch((74.5, 62.5), 14, 3.2, boxstyle="round,pad=0.1,rounding_size=0.6",
                            facecolor='#DBEAFE', edgecolor='#3B82F6', linewidth=0.8, zorder=5)
    ax.add_patch(badge3)
    ax.text(81.5, 64.1, "DEV_NODE", ha='center', va='center', fontsize=7.0, fontweight='bold', color='#1D4ED8', zorder=6)
    ax.text(81.5, 56.5, "步骤 3\n混匀搅拌", ha='center', va='center', fontsize=11, fontweight='bold', color='#0F172A', zorder=5)
    ax.text(81.5, 48.0, "执行能力: stir\n目标设定: 600rpm\n设备: 顶置搅拌器", ha='center', va='center', fontsize=7.5, color='#475569', zorder=5)

    # Port in Step 3
    p_in3 = FancyBboxPatch((74, 38), 15, 3.2, boxstyle="round,pad=0.1,rounding_size=0.4",
                           facecolor='#FFEDD5', edgecolor='#F97316', linewidth=0.8, zorder=5)
    ax.add_patch(p_in3)
    ax.text(81.5, 39.6, "PORT_IN: 产物状态", ha='center', va='center', fontsize=6.5, fontweight='bold', color='#C2410C', zorder=6)

    # END Terminal
    end_c = patches.Circle((95.5, 52), 3.5, facecolor='#64748B', edgecolor='#475569', linewidth=2, zorder=4)
    ax.add_patch(end_c)
    ax.text(95.5, 52, "END", ha='center', va='center', fontsize=6.8, fontweight='bold', color='#FFFFFF', zorder=5)

    # Control Flow Arrows (Interface Connections: ACTIVE signals)
    ax.add_patch(FancyArrowPatch((14.5, 52), (19.5, 52), arrowstyle='-|>', mutation_scale=12, linewidth=2.2, color='#2563EB', zorder=6))
    
    ax.add_patch(FancyArrowPatch((39.5, 53), (45.5, 53), arrowstyle='-|>', mutation_scale=13, linewidth=2.4, color='#2563EB', zorder=6))
    ax.text(42.5, 56.5, "ACTIVE\n控制激活", ha='center', va='center', fontsize=7.0, fontweight='bold', color='#1D4ED8', zorder=7)

    ax.add_patch(FancyArrowPatch((65.5, 53), (71.5, 53), arrowstyle='-|>', mutation_scale=13, linewidth=2.4, color='#2563EB', zorder=6))
    ax.text(68.5, 56.5, "ACTIVE\n控制激活", ha='center', va='center', fontsize=7.0, fontweight='bold', color='#1D4ED8', zorder=7)

    ax.add_patch(FancyArrowPatch((91.5, 52), (92.0, 52), arrowstyle='-|>', mutation_scale=11, linewidth=2.0, color='#64748B', zorder=6))

    # Data Port Connections (Port Connections: Material & Context Flow)
    port_arrow1 = FancyArrowPatch((37, 36), (49, 36),
                                  connectionstyle="arc3,rad=-0.5",
                                  arrowstyle='-|>', mutation_scale=13, linewidth=2.0,
                                  linestyle=':', color='#EA580C', zorder=6)
    ax.add_patch(port_arrow1)
    ax.text(43, 26.5, "端口单向拉取: 样品上下文", ha='center', va='center', fontsize=7.2, fontweight='bold', color='#C2410C', zorder=7)

    port_arrow2 = FancyArrowPatch((63, 36), (75, 36),
                                  connectionstyle="arc3,rad=-0.5",
                                  arrowstyle='-|>', mutation_scale=13, linewidth=2.0,
                                  linestyle=':', color='#EA580C', zorder=6)
    ax.add_patch(port_arrow2)
    ax.text(69, 26.5, "端口单向拉取: 物料温度数据", ha='center', va='center', fontsize=7.2, fontweight='bold', color='#C2410C', zorder=7)

    # Bottom Legend Banner
    leg_box = FancyBboxPatch((6, 3.5), 88, 12.0, boxstyle="round,pad=0.4,rounding_size=1.2",
                            facecolor='#F8FAFC', edgecolor='#E2E8F0', linewidth=1.2, zorder=3)
    ax.add_patch(leg_box)
    ax.text(50, 12.5, "正交双连接通道体系 (Orthogonal Connection Channels)", ha='center', va='center',
            fontsize=8.5, fontweight='bold', color='#334155', zorder=4)

    # Legend items
    ax.plot([10, 18], [7.2, 7.2], color='#2563EB', linewidth=2.4, zorder=5)
    ax.text(20, 7.2, "接口连接 (控制流): 信号活边触发状态推进 (ACTIVE / EMIT)", ha='left', va='center', fontsize=7.2, color='#1E293B', zorder=5)

    ax.plot([54, 62], [7.2, 7.2], color='#EA580C', linewidth=2.2, linestyle=':', zorder=5)
    ax.text(64, 7.2, "端口连接 (数据流): 下游主动拉取物料数据，源数据严格冻结", ha='left', va='center', fontsize=7.2, color='#1E293B', zorder=5)

    out_path = os.path.join(out_dir, "02_横向概念结构图.png")
    plt.tight_layout()
    plt.savefig(out_path, dpi=300, facecolor=fig.get_facecolor(), edgecolor='none')
    plt.close()
    print("Saved:", out_path)


# =============================================================
# 3. 切片示意图 (横向步骤组织 + 每步走纵向链路)
# =============================================================
def create_slice_diagram():
    fig, ax = plt.subplots(figsize=(8.2, 7.2), dpi=300)
    fig.patch.set_facecolor('#F8FAFC')
    ax.set_facecolor('#F8FAFC')
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis('off')

    # Outer Card Border
    outer_card = FancyBboxPatch((1.5, 1.5), 97, 97, boxstyle="round,pad=0.5,rounding_size=2.0",
                                facecolor='#FFFFFF', edgecolor='#CBD5E1', linewidth=1.5, zorder=1)
    ax.add_patch(outer_card)

    # Header Title
    ax.text(50, 94.2, "任务执行切片示意 (Execution Slice)", ha='center', va='center',
            fontsize=15.5, fontweight='bold', color='#0F172A', zorder=4)
    ax.text(50, 90.2, "切片 = 横向任务步骤编排 × 纵向不可旁路链路的正交融合", ha='center', va='center',
            fontsize=9.2, color='#64748B', zorder=4)

    # Slice Definition Formula Banner
    formula_box = FancyBboxPatch((6, 81.5), 88, 5.8, boxstyle="round,pad=0.3,rounding_size=1.0",
                                 facecolor='#EFF6FF', edgecolor='#BFDBFE', linewidth=1, zorder=2)
    ax.add_patch(formula_box)
    ax.text(50, 84.4, r"执行切片 $\mathcal{F} = \langle V_F, E_I, E_P \rangle$ ：横向时序拓扑，纵向步步贯通底层闭环",
            ha='center', va='center', fontsize=8.8, fontweight='bold', color='#1E40AF', zorder=3)

    # Row Header Badges on Left
    row_headers = [
        ("工作流层\n(时序编排)", 69.5, '#DBEAFE', '#1D4ED8'),
        ("状态机层\n(闭环控制)", 48.5, '#EDE9FE', '#6D28D9'),
        ("适配器层\n(协议转换)", 27.5, '#CCFBF1', '#0F766E'),
        ("物理设备层\n(执行终端)", 10.5, '#F1F5F9', '#334155')
    ]
    for text, y_pos, bg_c, txt_c in row_headers:
        lbl_box = FancyBboxPatch((2.8, y_pos - 4.5), 11.5, 9.0, boxstyle="round,pad=0.2,rounding_size=0.8",
                                 facecolor=bg_c, edgecolor=txt_c, linewidth=0.8, zorder=3)
        ax.add_patch(lbl_box)
        ax.text(8.55, y_pos, text, ha='center', va='center', fontsize=6.8, fontweight='bold', color=txt_c, zorder=4)

    # Step Columns
    cols_x = [27.5, 54.5, 81.5]
    step_titles = ["步骤 1: 称重加样", "步骤 2: 恒温反应", "步骤 3: 搅拌混匀"]
    sm_titles = ["加样机状态机", "反应釜状态机", "搅拌器状态机"]
    adapt_titles = ["加样机 Adapter", "反应釜 Adapter", "搅拌器 Adapter"]
    dev_titles = ["固体自动加样仪", "夹套式加热反应釜", "顶置高速搅拌机"]

    # Horizontal Flow connecting top nodes
    for i in range(2):
        x1 = cols_x[i] + 10.5
        x2 = cols_x[i+1] - 10.5
        ax.add_patch(FancyArrowPatch((x1, 69.5), (x2, 69.5), arrowstyle='-|>', mutation_scale=13,
                                     linewidth=2.4, color='#2563EB', zorder=8))
        ax.text((x1+x2)/2, 72.2, "ACTIVE", ha='center', va='center', fontsize=7.2, fontweight='bold', color='#1D4ED8', zorder=8)

    # Slice Projection Columns
    for idx, cx in enumerate(cols_x):
        # Column slice background container (translucent slice column)
        col_bg = FancyBboxPatch((cx - 11.5, 4.5), 23, 73, boxstyle="round,pad=0.3,rounding_size=1.5",
                                facecolor='#F0F9FF', edgecolor='#93C5FD', linewidth=1.2, linestyle=':', zorder=2)
        ax.add_patch(col_bg)

        # 1. Top: Workflow Node
        n_top = FancyBboxPatch((cx - 10, 64.5), 20, 10, boxstyle="round,pad=0.3,rounding_size=1.2",
                               facecolor='#1E3A8A', edgecolor='#3B82F6', linewidth=1.8, zorder=5)
        ax.add_patch(n_top)
        ax.text(cx, 70.8, step_titles[idx], ha='center', va='center',
                fontsize=8.2, fontweight='bold', color='#FFFFFF', zorder=6)
        ax.text(cx, 67.2, "[工作流执行节点]", ha='center', va='center',
                fontsize=7.0, color='#93C5FD', zorder=6)

        # Vertical Arrow 1: Node -> State Machine
        ax.add_patch(FancyArrowPatch((cx - 2.5, 64.5), (cx - 2.5, 54.0),
                                     arrowstyle='-|>', mutation_scale=11, linewidth=1.6, color='#2563EB', zorder=6))
        ax.add_patch(FancyArrowPatch((cx + 2.5, 54.0), (cx + 2.5, 64.5),
                                     arrowstyle='-|>', mutation_scale=11, linewidth=1.6, linestyle='--', color='#059669', zorder=6))

        # 2. Middle: State Machine (Gatekeeper)
        n_sm = FancyBboxPatch((cx - 10, 43.5), 20, 10.5, boxstyle="round,pad=0.3,rounding_size=1.2",
                              facecolor='#F5F3FF', edgecolor='#7C3AED', linewidth=2.0, zorder=5)
        ax.add_patch(n_sm)
        ax.text(cx, 50.0, sm_titles[idx], ha='center', va='center',
                fontsize=8.2, fontweight='bold', color='#5B21B6', zorder=6)
        ax.text(cx, 46.2, "★ 独占发令中枢", ha='center', va='center',
                fontsize=7.0, fontweight='bold', color='#7C3AED', zorder=6)

        # Vertical Arrow 2: State Machine -> Adapter
        ax.add_patch(FancyArrowPatch((cx - 2.5, 43.5), (cx - 2.5, 32.5),
                                     arrowstyle='-|>', mutation_scale=11, linewidth=1.6, color='#7C3AED', zorder=6))
        ax.add_patch(FancyArrowPatch((cx + 2.5, 32.5), (cx + 2.5, 43.5),
                                     arrowstyle='-|>', mutation_scale=11, linewidth=1.6, linestyle='--', color='#059669', zorder=6))

        # 3. Adapter Layer
        n_ad = FancyBboxPatch((cx - 10, 23.5), 20, 8.8, boxstyle="round,pad=0.3,rounding_size=1.0",
                              facecolor='#0F766E', edgecolor='#14B8A6', linewidth=1.5, zorder=5)
        ax.add_patch(n_ad)
        ax.text(cx, 27.9, adapt_titles[idx], ha='center', va='center',
                fontsize=8.0, fontweight='bold', color='#FFFFFF', zorder=6)

        # Vertical Arrow 3: Adapter -> Hardware
        ax.add_patch(FancyArrowPatch((cx, 23.5), (cx, 15.5),
                                     arrowstyle='-|>', mutation_scale=10, linewidth=1.5, color='#0D9488', zorder=6))

        # 4. Hardware Layer
        n_hw = FancyBboxPatch((cx - 10, 6.5), 20, 8.8, boxstyle="round,pad=0.3,rounding_size=1.0",
                              facecolor='#1E293B', edgecolor='#475569', linewidth=1.5, zorder=5)
        ax.add_patch(n_hw)
        ax.text(cx, 10.9, dev_titles[idx], ha='center', va='center',
                fontsize=8.0, fontweight='bold', color='#F8FAFC', zorder=6)

    out_path = os.path.join(out_dir, "03_切片示意图.png")
    plt.tight_layout()
    plt.savefig(out_path, dpi=300, facecolor=fig.get_facecolor(), edgecolor='none')
    plt.close()
    print("Saved:", out_path)


# =============================================================
# 4. 三图全景组合图 (Panorama 3-in-1 for Slide 6)
# =============================================================
def create_panorama_diagram():
    fig, axes = plt.subplots(1, 3, figsize=(18, 6.2), dpi=300)
    fig.patch.set_facecolor('#F8FAFC')
    
    import matplotlib.image as mpimg
    img1 = mpimg.imread(os.path.join(out_dir, "01_纵向概念结构图.png"))
    img2 = mpimg.imread(os.path.join(out_dir, "02_横向概念结构图.png"))
    img3 = mpimg.imread(os.path.join(out_dir, "03_切片示意图.png"))

    axes[0].imshow(img1)
    axes[0].axis('off')
    axes[0].set_title("① 纵向概念结构图 (单步闭环)", fontsize=13.5, fontweight='bold', color='#0F172A', pad=10)

    axes[1].imshow(img2)
    axes[1].axis('off')
    axes[1].set_title("② 横向概念结构图 (多步编排)", fontsize=13.5, fontweight='bold', color='#0F172A', pad=10)

    axes[2].imshow(img3)
    axes[2].axis('off')
    axes[2].set_title("③ 任务执行切片 (正交融合)", fontsize=13.5, fontweight='bold', color='#0F172A', pad=10)

    plt.tight_layout()
    out_path = os.path.join(out_dir, "00_执行链路与任务切片_三图全景.png")
    plt.savefig(out_path, dpi=300, facecolor=fig.get_facecolor(), edgecolor='none')
    plt.close()
    print("Saved Panorama:", out_path)


# =============================================================
# 5. 更新 PPTX Slide 6
# =============================================================
def update_pptx():
    pptx_files = glob.glob(r"d:\SmartLab2.0\doc\*汇报.pptx")
    if not pptx_files:
        print("No pptx found")
        return
    src_pptx = pptx_files[0]
    prs = Presentation(src_pptx)
    slide = prs.slides[5] # Slide 6
    
    img1_path = os.path.join(out_dir, "01_纵向概念结构图.png")
    img2_path = os.path.join(out_dir, "02_横向概念结构图.png")
    img3_path = os.path.join(out_dir, "03_切片示意图.png")

    # Card 6: left: 365760, top: 2468880, width: 3657600, height: 3200400
    # Card 7: left: 4251960, top: 2468880, width: 3657600, height: 3200400
    # Card 8: left: 8138160, top: 2468880, width: 3657600, height: 3200400
    cards_info = [
        ("Rounded Rectangle 6", img1_path),
        ("Rounded Rectangle 7", img2_path),
        ("Rounded Rectangle 8", img3_path)
    ]
    
    for shp_name, img_path in cards_info:
        for shp in slide.shapes:
            if shp.name == shp_name:
                left = shp.left
                top = shp.top
                width = shp.width
                height = shp.height
                # Clear placeholder text
                if shp.has_text_frame:
                    shp.text_frame.text = ""
                # Add picture on top of card with slight padding
                pad = 40000
                slide.shapes.add_picture(img_path, left + pad, top + pad, width - 2*pad, height - 2*pad)
                break

    dst_pptx = r"d:\SmartLab2.0\doc\EUROSIM2026_于涛_执行链路汇报_含图版.pptx"
    prs.save(dst_pptx)
    print("Updated PPTX saved to:", dst_pptx)

if __name__ == "__main__":
    create_vertical_diagram()
    create_horizontal_diagram()
    create_slice_diagram()
    create_panorama_diagram()
    update_pptx()
